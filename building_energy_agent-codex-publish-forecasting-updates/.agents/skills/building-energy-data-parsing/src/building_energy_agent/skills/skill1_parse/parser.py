"""Memory-bounded XLSX to DatasetBundle parser for the V0.7 reference shape."""

from __future__ import annotations

import hashlib
import json
import math
import resource
import shutil
import tempfile
import time
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from building_energy_agent.contracts.bundle_layouts import DATASET_BUNDLE_V07
from building_energy_agent.contracts.fingerprints import dataset_schema_fingerprint
from building_energy_agent.contracts.skill1_v07 import (
    DatasetBundleManifest,
    PointCatalogEntry,
    Skill1InlineSummary,
)
from building_energy_agent.core.errors import SkillNeedsInputError, SkillParseError
from building_energy_agent.skills.skill1_parse.sources import (
    RegisteredSource,
    sha256_file,
)

PARSE_RULE_VERSION = "xlsx-two-header-v1"
PARQUET_BATCH_RECORDS = 50_000
TIMESERIES_SCHEMA = pa.schema(
    [
        pa.field("timestamp", pa.timestamp("us", tz="UTC"), nullable=False),
        pa.field("point_id", pa.string(), nullable=False),
        pa.field("value_raw", pa.float64(), nullable=True),
    ]
)
POINTS_SCHEMA = pa.schema(
    [
        pa.field("point_id", pa.string(), nullable=False),
        pa.field("name", pa.string(), nullable=False),
        pa.field("description", pa.string(), nullable=False),
        pa.field("unit", pa.string(), nullable=True),
        pa.field("variable_type", pa.string(), nullable=False),
        pa.field("expected_interval_seconds", pa.int64(), nullable=False),
    ]
)


@dataclass(frozen=True, slots=True)
class ParseOutcome:
    """Bundle identity, location, and lightweight API summary."""

    dataset_bundle_id: str
    bundle_path: Path
    summary: Skill1InlineSummary
    report_summary: dict[str, Any]


@dataclass(frozen=True, slots=True)
class _ColumnMetadata:
    column_index: int
    original_column_name: str
    point_id: str
    name: str
    description: str
    unit: str | None
    variable_type: str

    @property
    def conflict_key(self) -> tuple[str, str, str | None, str]:
        return (self.name, self.description, self.unit, self.variable_type)


def _json_write(path: Path, value: object) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _canonical_id(sources: list[RegisteredSource], source_timezone: str | None) -> str:
    identity = {
        "parse_rule_version": PARSE_RULE_VERSION,
        "source_timezone": source_timezone,
        "source_sha256": [source.sha256 for source in sources],
    }
    encoded = json.dumps(
        identity, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    digest = hashlib.sha256(encoded).hexdigest()
    return f"dataset-{digest[:20]}-{PARSE_RULE_VERSION}"


def _partition_name(source_index: int, sheet_index: int, sheet_name: str) -> str:
    digest = hashlib.sha256(sheet_name.encode("utf-8")).hexdigest()[:8]
    return f"source-{source_index:02d}-sheet-{sheet_index:03d}-{digest}"


def _description_parts(raw_description: object) -> tuple[str, str, str | None]:
    if raw_description is None:
        return "", "", None
    text = str(raw_description).strip().replace("｜", "|")
    parts = [part.strip() for part in text.split("|")]
    if len(parts) >= 3:
        return parts[0], "|".join(parts[1:-1]), parts[-1] or None
    if len(parts) == 2:
        return parts[0], parts[1], None
    return "", text, None


def _timestamp_from_value(
    value: object,
    *,
    source_zone: ZoneInfo | None,
    workbook_epoch: datetime,
) -> datetime:
    timestamp: datetime
    if isinstance(value, datetime):
        timestamp = value
    elif isinstance(value, date):
        timestamp = datetime.combine(value, datetime.min.time())
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        timestamp = from_excel(value, workbook_epoch)
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            raise ValueError("empty timestamp")
        normalized = text[:-1] + "+00:00" if text.endswith("Z") else text
        try:
            timestamp = datetime.fromisoformat(normalized)
        except ValueError:
            for format_text in (
                "%Y/%m/%d %H:%M:%S",
                "%Y-%m-%d %H:%M:%S",
                "%Y/%m/%d %H:%M",
                "%Y-%m-%d %H:%M",
            ):
                try:
                    timestamp = datetime.strptime(text, format_text)
                    break
                except ValueError:
                    continue
            else:
                raise ValueError(f"unparseable timestamp: {text[:80]!r}")
    else:
        raise ValueError(f"unsupported timestamp type: {type(value).__name__}")

    if timestamp.tzinfo is None or timestamp.utcoffset() is None:
        if source_zone is None:
            raise SkillNeedsInputError(
                "source timestamps are timezone-naive; source_timezone is required"
            )
        timestamp = timestamp.replace(tzinfo=source_zone)
    return timestamp.astimezone(UTC)


def _encoded_text_value(text: str) -> float:
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    return float(int(digest[:13], 16))


def _numeric_value(
    value: object,
    *,
    variable_type: str,
) -> tuple[float | None, tuple[str, float] | None, bool]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None, False
    if isinstance(value, bool):
        return float(int(value)), (str(value), float(int(value))), False
    if isinstance(value, (int, float, Decimal)):
        number = float(value)
        if math.isfinite(number):
            return number, None, False
        return None, None, True
    if isinstance(value, str):
        text = value.strip()
        state_tokens = {
            "off": 0.0,
            "false": 0.0,
            "closed": 0.0,
            "停止": 0.0,
            "关闭": 0.0,
            "关": 0.0,
            "on": 1.0,
            "true": 1.0,
            "open": 1.0,
            "运行": 1.0,
            "开启": 1.0,
            "开": 1.0,
        }
        lowered = text.casefold()
        if lowered in state_tokens:
            number = state_tokens[lowered]
            return number, (text, number), False
        try:
            number = float(Decimal(text))
        except (InvalidOperation, ValueError, OverflowError):
            if variable_type in {"state", "enum"}:
                number = _encoded_text_value(text)
                return number, (text, number), False
            return None, None, True
        if math.isfinite(number):
            return number, None, False
    return None, None, True


def _counter_median(counter: Counter[int]) -> int | None:
    total = sum(counter.values())
    if total == 0:
        return None
    positions = ((total - 1) // 2, total // 2)
    selected: list[int] = []
    cumulative = 0
    for value, count in sorted(counter.items()):
        previous = cumulative
        cumulative += count
        for position in positions[len(selected) :]:
            if previous <= position < cumulative:
                selected.append(value)
            else:
                break
        if len(selected) == 2:
            break
    return max(1, round(sum(selected) / len(selected)))


def _max_rss_bytes() -> int:
    usage = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return int(usage * 1024)


def _contains_vba(path: Path) -> bool:
    if path.suffix.casefold() != ".xlsm":
        return False
    try:
        with zipfile.ZipFile(path) as archive:
            return any(name.casefold().endswith("vbaproject.bin") for name in archive.namelist())
    except zipfile.BadZipFile as exc:
        raise SkillParseError("invalid XLSM container") from exc


class XlsxDatasetParser:
    """Parse two-header XLSX worksheets into partitioned long-form Parquet."""

    def __init__(self, output_root: str | Path) -> None:
        self.output_root = Path(output_root).expanduser().resolve()

    def parse(
        self,
        sources: list[RegisteredSource],
        *,
        source_timezone: str | None,
        strict_mode: bool,
    ) -> ParseOutcome:
        if not sources:
            raise SkillParseError("at least one registered source is required")
        source_zone: ZoneInfo | None = None
        if source_timezone:
            try:
                source_zone = ZoneInfo(source_timezone)
            except ZoneInfoNotFoundError as exc:
                raise SkillNeedsInputError(
                    f"unknown source_timezone: {source_timezone!r}"
                ) from exc

        bundle_id = _canonical_id(sources, source_timezone)
        final_root = self.output_root / bundle_id
        if final_root.exists():
            return self._load_existing(final_root, bundle_id)

        self.output_root.mkdir(parents=True, exist_ok=True)
        temporary_root = Path(
            tempfile.mkdtemp(prefix=f".{bundle_id}-", dir=self.output_root)
        )
        try:
            outcome = self._build(
                temporary_root,
                bundle_id,
                sources,
                source_zone=source_zone,
                source_timezone=source_timezone,
                strict_mode=strict_mode,
            )
            DATASET_BUNDLE_V07.validate(temporary_root)
            temporary_root.rename(final_root)
            return ParseOutcome(
                dataset_bundle_id=outcome.dataset_bundle_id,
                bundle_path=final_root,
                summary=outcome.summary,
                report_summary=outcome.report_summary,
            )
        except Exception:
            shutil.rmtree(temporary_root, ignore_errors=True)
            raise

    def _build(
        self,
        root: Path,
        bundle_id: str,
        sources: list[RegisteredSource],
        *,
        source_zone: ZoneInfo | None,
        source_timezone: str | None,
        strict_mode: bool,
    ) -> ParseOutcome:
        started = time.monotonic()
        (root / "data" / "timeseries").mkdir(parents=True)
        (root / "data" / "tables").mkdir(parents=True)

        point_metadata: dict[str, _ColumnMetadata] = {}
        point_differences: dict[str, Counter[int]] = defaultdict(Counter)
        state_encodings: dict[str, dict[str, float]] = defaultdict(dict)
        worksheet_reports: list[dict[str, Any]] = []
        global_warnings: list[str] = []
        output_file_hashes: dict[str, str] = {}
        global_start: datetime | None = None
        global_end: datetime | None = None
        total_records = 0
        total_nulls = 0
        successful_files = 0
        failed_files = 0
        successful_sheets = 0
        failed_sheets = 0

        for source_index, source in enumerate(sources, start=1):
            if source.path.suffix.lower() not in {".xlsx", ".xlsm"}:
                failed_files += 1
                message = (
                    f"unsupported source format for reference parser: "
                    f"{source.path.suffix.lower()}"
                )
                if strict_mode:
                    raise SkillParseError(message)
                global_warnings.append(message)
                continue
            workbook = None
            formula_workbook = None
            source_had_success = False
            try:
                workbook = load_workbook(
                    source.path,
                    read_only=True,
                    data_only=True,
                    keep_links=False,
                    keep_vba=False,
                )
                macro_present = _contains_vba(source.path)
                if source.path.suffix.casefold() == ".xlsm":
                    formula_workbook = load_workbook(
                        source.path,
                        read_only=True,
                        data_only=False,
                        keep_links=False,
                        keep_vba=False,
                    )
                for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
                    formula_worksheet = None
                    if formula_workbook is not None:
                        formula_worksheet = formula_workbook[worksheet.title]
                    metadata_snapshot = point_metadata.copy()
                    differences_snapshot = {
                        key: value.copy() for key, value in point_differences.items()
                    }
                    encodings_snapshot = {
                        key: value.copy() for key, value in state_encodings.items()
                    }
                    try:
                        report = self._parse_worksheet(
                            root,
                            source,
                            source_index,
                            worksheet,
                            sheet_index,
                            workbook.epoch,
                            source_zone,
                            point_metadata,
                            point_differences,
                            state_encodings,
                            formula_worksheet=formula_worksheet,
                            macro_present=macro_present,
                        )
                    except SkillNeedsInputError:
                        point_metadata.clear()
                        point_metadata.update(metadata_snapshot)
                        point_differences.clear()
                        point_differences.update(differences_snapshot)
                        state_encodings.clear()
                        state_encodings.update(encodings_snapshot)
                        raise
                    except Exception as exc:
                        failed_sheets += 1
                        point_metadata.clear()
                        point_metadata.update(metadata_snapshot)
                        point_differences.clear()
                        point_differences.update(differences_snapshot)
                        state_encodings.clear()
                        state_encodings.update(encodings_snapshot)
                        message = f"worksheet {worksheet.title!r} skipped: {exc}"
                        worksheet_reports.append(
                            {
                                "source_ref": source.source_ref,
                                "worksheet_name": worksheet.title,
                                "status": "failed",
                                "warnings": [message],
                            }
                        )
                        if strict_mode:
                            raise SkillParseError(message) from exc
                        global_warnings.append(message)
                        continue

                    source_had_success = True
                    successful_sheets += 1
                    worksheet_reports.append(report)
                    global_warnings.extend(report.get("warnings", []))
                    total_records += report["record_count"]
                    total_nulls += report["null_count"]
                    report_start = datetime.fromisoformat(report["time_start"])
                    report_end = datetime.fromisoformat(report["time_end"])
                    global_start = (
                        report_start
                        if global_start is None
                        else min(global_start, report_start)
                    )
                    global_end = (
                        report_end if global_end is None else max(global_end, report_end)
                    )
                    output_file_hashes[report["output_partition"]] = report[
                        "output_sha256"
                    ]
                if source_had_success:
                    successful_files += 1
                else:
                    failed_files += 1
            except SkillNeedsInputError:
                raise
            except Exception as exc:
                if strict_mode or isinstance(exc, SkillParseError):
                    raise
                failed_files += int(not source_had_success)
                global_warnings.append(
                    f"source {source.original_filename!r} parse failure: {exc}"
                )
            finally:
                if workbook is not None:
                    workbook.close()
                if formula_workbook is not None:
                    formula_workbook.close()

        if not worksheet_reports or successful_sheets == 0:
            raise SkillParseError("no usable time-series worksheets were parsed")
        if global_start is None or global_end is None:
            raise SkillParseError("parsed worksheets did not contain valid timestamps")

        point_rows: list[dict[str, Any]] = []
        for point_id, metadata in sorted(point_metadata.items()):
            interval = _counter_median(point_differences[point_id])
            if interval is None:
                raise SkillNeedsInputError(
                    f"point {point_id!r} has insufficient timestamps to infer interval"
                )
            PointCatalogEntry(
                point_id=point_id,
                name=metadata.name,
                description=metadata.description,
                unit=metadata.unit,
                variable_type=metadata.variable_type,
                expected_interval_seconds=interval,
            )
            differences = point_differences[point_id]
            irregular = sum(
                count
                for difference, count in differences.items()
                if difference % interval != 0
            )
            difference_count = sum(differences.values())
            if difference_count >= 10 and irregular / difference_count > 0.05:
                global_warnings.append(
                    f"point {point_id!r} has visibly irregular sampling "
                    f"({irregular}/{difference_count} positive intervals)"
                )
            point_rows.append(
                {
                    "point_id": point_id,
                    "name": metadata.name,
                    "description": metadata.description,
                    "unit": metadata.unit,
                    "variable_type": metadata.variable_type,
                    "expected_interval_seconds": interval,
                }
            )

        points_path = root / "points.parquet"
        pq.write_table(
            pa.Table.from_pylist(point_rows, schema=POINTS_SCHEMA),
            points_path,
            compression="zstd",
        )
        fingerprint = dataset_schema_fingerprint(point_rows)
        manifest = DatasetBundleManifest(
            dataset_bundle_id=bundle_id,
            schema_fingerprint=fingerprint,
            time_start=global_start,
            time_end=global_end,
        )
        manifest_payload = manifest.model_dump(mode="json")
        manifest_payload["time_start"] = manifest.time_start.isoformat()
        manifest_payload["time_end"] = manifest.time_end.isoformat()
        _json_write(root / "manifest.json", manifest_payload)
        output_file_hashes["manifest.json"] = sha256_file(root / "manifest.json")
        output_file_hashes["points.parquet"] = sha256_file(points_path)

        source_reports = [
            {
                "source_ref": source.source_ref,
                "original_filename": source.original_filename,
                "size_bytes": source.size_bytes,
                "sha256": source.sha256,
            }
            for source in sources
        ]
        report: dict[str, Any] = {
            "dataset_bundle_id": bundle_id,
            "parse_rule_version": PARSE_RULE_VERSION,
            "source_timezone": source_timezone,
            "sources": source_reports,
            "worksheet_count": len(worksheet_reports),
            "successful_worksheet_count": successful_sheets,
            "failed_worksheet_count": failed_sheets,
            "point_count": len(point_rows),
            "record_count": total_records,
            "null_count": total_nulls,
            "time_start": global_start.isoformat(),
            "time_end": global_end.isoformat(),
            "worksheets": worksheet_reports,
            "state_encodings": state_encodings,
            "warnings": global_warnings,
            "output_file_hashes": output_file_hashes,
            "elapsed_seconds": round(time.monotonic() - started, 3),
            "peak_memory_bytes_approx": _max_rss_bytes(),
            "output_size_bytes": 0,
        }
        report_path = root / "parse_report.json"
        _json_write(report_path, report)
        report["output_size_bytes"] = sum(
            path.stat().st_size for path in root.rglob("*") if path.is_file()
        )
        _json_write(report_path, report)

        summary = Skill1InlineSummary(
            successful_file_count=successful_files,
            failed_file_count=failed_files,
            successful_worksheet_count=successful_sheets,
            failed_worksheet_count=failed_sheets,
            timeseries_table_count=successful_sheets,
            regular_table_count=0,
            point_count=len(point_rows),
            record_count=total_records,
            time_range={
                "time_start": global_start.isoformat(),
                "time_end": global_end.isoformat(),
            },
            main_warnings=global_warnings[:20],
            needs_input=False,
        )
        return ParseOutcome(bundle_id, root, summary, report)

    def _parse_worksheet(
        self,
        root: Path,
        source: RegisteredSource,
        source_index: int,
        worksheet: Any,
        sheet_index: int,
        workbook_epoch: datetime,
        source_zone: ZoneInfo | None,
        point_metadata: dict[str, _ColumnMetadata],
        point_differences: dict[str, Counter[int]],
        state_encodings: dict[str, dict[str, float]],
        *,
        formula_worksheet: Any | None = None,
        macro_present: bool = False,
    ) -> dict[str, Any]:
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        formula_rows = None
        if formula_worksheet is not None:
            formula_worksheet.reset_dimensions()
            formula_rows = formula_worksheet.iter_rows()
        try:
            identifiers = tuple(next(rows))
            descriptions = tuple(next(rows))
            if formula_rows is not None:
                next(formula_rows)
                next(formula_rows)
        except StopIteration as exc:
            raise SkillParseError("worksheet does not contain two header rows") from exc
        if not identifiers:
            raise SkillParseError("worksheet point-id header is empty")
        timestamp_header = str(identifiers[0] or "").strip().casefold()
        if timestamp_header not in {
            "time",
            "timestamp",
            "datetime",
            "时间",
            "日期时间",
        }:
            raise SkillNeedsInputError(
                f"worksheet {worksheet.title!r} has ambiguous timestamp column "
                f"{str(identifiers[0])[:80]!r}"
            )

        variable_type = "state" if "状态" in worksheet.title else "continuous"
        column_count = max(len(identifiers), len(descriptions))
        columns: list[_ColumnMetadata] = []
        column_warnings: list[str] = []
        for column_index in range(1, column_count):
            raw_identifier = (
                identifiers[column_index] if column_index < len(identifiers) else None
            )
            raw_description = (
                descriptions[column_index]
                if column_index < len(descriptions)
                else None
            )
            if raw_identifier is None and raw_description is None:
                continue
            if raw_identifier is None or not str(raw_identifier).strip():
                point_id = (
                    f"generated-{source.sha256[:10]}-{sheet_index:03d}-"
                    f"column-{column_index + 1:04d}"
                )
                column_warnings.append(
                    f"column {column_index + 1} had no point ID; used {point_id!r}"
                )
                original_column_name = ""
            else:
                original_column_name = str(raw_identifier).strip()
                point_id = original_column_name
            name, description, unit = _description_parts(raw_description)
            metadata = _ColumnMetadata(
                column_index=column_index,
                original_column_name=original_column_name,
                point_id=point_id,
                name=name,
                description=description,
                unit=unit,
                variable_type=variable_type,
            )
            existing = point_metadata.get(point_id)
            if existing is not None and existing.conflict_key != metadata.conflict_key:
                raise SkillNeedsInputError(
                    f"point_id {point_id!r} has conflicting metadata: "
                    f"{existing.conflict_key!r} versus {metadata.conflict_key!r}"
                )
            point_metadata.setdefault(point_id, metadata)
            columns.append(metadata)
        if not columns:
            raise SkillParseError("worksheet does not contain point columns")

        partition = _partition_name(source_index, sheet_index, worksheet.title)
        relative_output = f"data/timeseries/sheet={partition}/data.parquet"
        output_path = root / relative_output
        output_path.parent.mkdir(parents=True)
        writer = pq.ParquetWriter(
            output_path,
            TIMESERIES_SCHEMA,
            compression="zstd",
            use_dictionary=["point_id"],
        )
        timestamps: list[datetime] = []
        point_ids: list[str] = []
        values: list[float | None] = []
        sheet_last_valid: dict[str, datetime] = {}
        sheet_timestamp_differences: Counter[int] = Counter()
        previous_sheet_timestamp: datetime | None = None
        row_count = 0
        timestamp_row_count = 0
        null_count = 0
        invalid_value_count = 0
        missing_formula_cache_count = 0
        invalid_timestamp_count = 0
        time_start: datetime | None = None
        time_end: datetime | None = None

        def flush() -> None:
            if not timestamps:
                return
            table = pa.Table.from_arrays(
                [
                    pa.array(timestamps, type=TIMESERIES_SCHEMA.field("timestamp").type),
                    pa.array(point_ids, type=pa.string()),
                    pa.array(values, type=pa.float64()),
                ],
                schema=TIMESERIES_SCHEMA,
            )
            writer.write_table(table)
            timestamps.clear()
            point_ids.clear()
            values.clear()

        try:
            for row in rows:
                formula_row = next(formula_rows, ()) if formula_rows is not None else ()
                timestamp_value = row[0] if row else None
                if timestamp_value is None:
                    continue
                try:
                    timestamp = _timestamp_from_value(
                        timestamp_value,
                        source_zone=source_zone,
                        workbook_epoch=workbook_epoch,
                    )
                except SkillNeedsInputError:
                    raise
                except ValueError:
                    invalid_timestamp_count += 1
                    continue
                timestamp_row_count += 1
                time_start = timestamp if time_start is None else min(time_start, timestamp)
                time_end = timestamp if time_end is None else max(time_end, timestamp)
                if previous_sheet_timestamp is not None:
                    difference = round(
                        (timestamp - previous_sheet_timestamp).total_seconds()
                    )
                    if difference > 0:
                        sheet_timestamp_differences[difference] += 1
                previous_sheet_timestamp = timestamp

                for column in columns:
                    raw_value = (
                        row[column.column_index]
                        if column.column_index < len(row)
                        else None
                    )
                    if (
                        raw_value is None
                        and column.column_index < len(formula_row)
                        and formula_row[column.column_index].data_type == "f"
                    ):
                        missing_formula_cache_count += 1
                    number, encoding, invalid = _numeric_value(
                        raw_value, variable_type=column.variable_type
                    )
                    if invalid:
                        invalid_value_count += 1
                    if number is None:
                        null_count += 1
                    else:
                        previous = sheet_last_valid.get(column.point_id)
                        if previous is not None:
                            difference = round((timestamp - previous).total_seconds())
                            if difference > 0:
                                point_differences[column.point_id][difference] += 1
                        sheet_last_valid[column.point_id] = timestamp
                    if encoding is not None:
                        text, encoded_value = encoding
                        prior = state_encodings[column.point_id].get(text)
                        if prior is not None and prior != encoded_value:
                            raise SkillParseError(
                                f"non-deterministic state encoding for {column.point_id!r}"
                            )
                        state_encodings[column.point_id][text] = encoded_value
                    timestamps.append(timestamp)
                    point_ids.append(column.point_id)
                    values.append(number)
                    row_count += 1
                    if len(timestamps) >= PARQUET_BATCH_RECORDS:
                        flush()
            flush()
        except Exception:
            writer.close()
            output_path.unlink(missing_ok=True)
            raise
        else:
            writer.close()

        if timestamp_row_count == 0 or time_start is None or time_end is None:
            output_path.unlink(missing_ok=True)
            raise SkillParseError("worksheet contains no valid timestamp rows")
        sheet_interval = _counter_median(sheet_timestamp_differences)
        if sheet_interval is None:
            raise SkillNeedsInputError(
                f"worksheet {worksheet.title!r} has insufficient timestamps"
            )
        for column in columns:
            if not point_differences[column.point_id]:
                point_differences[column.point_id][sheet_interval] += 1
                column_warnings.append(
                    f"point {column.point_id!r} used worksheet interval fallback"
                )
        if invalid_timestamp_count:
            column_warnings.append(
                f"skipped {invalid_timestamp_count} rows with invalid timestamps"
            )
        if invalid_value_count:
            column_warnings.append(
                f"converted {invalid_value_count} non-numeric continuous values to null"
            )
        if missing_formula_cache_count:
            column_warnings.append(
                f"{missing_formula_cache_count} formula cells had no cached value; preserved as null"
            )
        if macro_present:
            column_warnings.append("XLSM VBA content was ignored and was not copied to the DatasetBundle")

        return {
            "source_ref": source.source_ref,
            "worksheet_name": worksheet.title,
            "status": "succeeded",
            "output_partition": relative_output,
            "output_sha256": sha256_file(output_path),
            "columns": [
                {
                    "original_column_name": column.original_column_name,
                    "point_id": column.point_id,
                    "name": column.name,
                    "description": column.description,
                    "unit": column.unit,
                    "variable_type": column.variable_type,
                }
                for column in columns
            ],
            "timestamp_row_count": timestamp_row_count,
            "record_count": row_count,
            "null_count": null_count,
            "time_start": time_start.isoformat(),
            "time_end": time_end.isoformat(),
            "warnings": column_warnings,
        }

    def _load_existing(self, root: Path, bundle_id: str) -> ParseOutcome:
        DATASET_BUNDLE_V07.validate(root)
        manifest = DatasetBundleManifest.model_validate_json(
            (root / "manifest.json").read_text(encoding="utf-8")
        )
        if manifest.dataset_bundle_id != bundle_id or root.name != bundle_id:
            raise SkillParseError("existing DatasetBundle identity is inconsistent")
        report = json.loads((root / "parse_report.json").read_text(encoding="utf-8"))
        summary = Skill1InlineSummary(
            successful_file_count=sum(
                1 for source in report.get("sources", []) if source
            ),
            failed_file_count=0,
            successful_worksheet_count=report["successful_worksheet_count"],
            failed_worksheet_count=report["failed_worksheet_count"],
            timeseries_table_count=report["successful_worksheet_count"],
            regular_table_count=0,
            point_count=report["point_count"],
            record_count=report["record_count"],
            time_range={
                "time_start": report["time_start"],
                "time_end": report["time_end"],
            },
            main_warnings=report.get("warnings", [])[:20],
            needs_input=False,
        )
        return ParseOutcome(bundle_id, root, summary, report)
