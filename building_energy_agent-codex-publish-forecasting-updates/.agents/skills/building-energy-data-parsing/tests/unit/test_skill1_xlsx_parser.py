from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pyarrow.parquet as pq
from openpyxl import Workbook

from building_energy_agent.contracts.skill1_v07 import ParseRequest, SourceReference
from building_energy_agent.skills.skill1_parse.artifacts import LocalDatasetBundleStore
from building_energy_agent.skills.skill1_parse.service import Skill1ParseService
from building_energy_agent.skills.skill1_parse.sources import LocalSourceRegistry
from building_energy_agent.skills.skill1_parse.validation import validate_dataset_bundle


def _write_reference_workbook(path: Path) -> None:
    workbook = Workbook()
    load_sheet = workbook.active
    load_sheet.title = "总冷机冷负荷"
    load_sheet.append(["time", "P-LOAD"])
    load_sheet.append(["点位描述", "冷机系统|总冷机冷负荷|kW"])
    load_sheet.append([datetime(2026, 1, 1, 8, 0), 10.0])
    load_sheet.append([datetime(2026, 1, 1, 8, 1), None])
    load_sheet.append([datetime(2026, 1, 1, 8, 2), 12.0])

    state_sheet = workbook.create_sheet("冷机启停状态")
    state_sheet.append(["time", "P-STATE"])
    state_sheet.append(["点位描述", "冷机1|启停状态|"])
    state_sheet.append([datetime(2026, 1, 1, 8, 0), "关闭"])
    state_sheet.append([datetime(2026, 1, 1, 8, 1), "开启"])
    state_sheet.append([datetime(2026, 1, 1, 8, 2), "关闭"])
    workbook.save(path)


def _service_for(source_path: Path, artifact_root: Path) -> tuple[
    Skill1ParseService, SourceReference
]:
    registry = LocalSourceRegistry()
    registered = registry.register(source_path)
    return (
        Skill1ParseService(registry, artifact_root),
        SourceReference(source_ref=registered.source_ref),
    )


def test_streaming_two_header_multi_sheet_xlsx_to_dataset_bundle(
    tmp_path: Path,
) -> None:
    source_path = tmp_path / "synthetic.xlsx"
    _write_reference_workbook(source_path)
    service, source = _service_for(source_path, tmp_path / "artifacts")

    started = service.start(
        ParseRequest(
            sources=[source],
            parse_hints={"source_timezone": "Asia/Shanghai"},
        )
    )
    assert started.status == "succeeded"
    result = service.result(started.job_id)
    bundle_path = (
        tmp_path / "artifacts" / "dataset-bundles" / result.dataset_bundle_id
    )
    manifest = json.loads((bundle_path / "manifest.json").read_text())
    artifact_store = LocalDatasetBundleStore(tmp_path / "artifacts")
    manifest_ref = artifact_store.manifest_ref(result.dataset_bundle_id)
    assert manifest_ref.bundle_id == result.dataset_bundle_id
    assert Path(manifest_ref.location_ref) == bundle_path / "manifest.json"

    visualization_manifests = list(
        (tmp_path / "artifacts" / "visualization-bundles").glob("*/manifest.json")
    )
    assert len(visualization_manifests) == 1
    visualization_manifest = json.loads(visualization_manifests[0].read_text())
    assert visualization_manifest["source"]["bundle_id"] == result.dataset_bundle_id

    assert bundle_path.name == result.dataset_bundle_id
    assert manifest["dataset_bundle_id"] == result.dataset_bundle_id
    assert manifest["time_start"] == "2026-01-01T00:00:00+00:00"
    assert manifest["time_end"] == "2026-01-01T00:02:00+00:00"
    assert set(manifest) == {
        "bundle_type",
        "bundle_version",
        "dataset_bundle_id",
        "schema_fingerprint",
        "time_start",
        "time_end",
    }

    points = pq.read_table(bundle_path / "points.parquet").to_pylist()
    assert {point["point_id"] for point in points} == {"P-LOAD", "P-STATE"}
    assert len({point["point_id"] for point in points}) == len(points)
    assert all(point["expected_interval_seconds"] > 0 for point in points)
    state_point = next(point for point in points if point["point_id"] == "P-STATE")
    assert state_point["variable_type"] == "state"

    parse_report = json.loads(
        (bundle_path / "parse_report.json").read_text(encoding="utf-8")
    )
    state_partition = next(
        entry
        for entry in parse_report["worksheets"]
        if entry["worksheet_name"] == "冷机启停状态"
    )
    state_values = pq.read_table(bundle_path / state_partition["output_partition"])[
        "value_raw"
    ].to_pylist()
    assert state_values == [0.0, 1.0, 0.0]
    assert parse_report["state_encodings"]["P-STATE"] == {
        "关闭": 0.0,
        "开启": 1.0,
    }

    validation = validate_dataset_bundle(bundle_path)
    assert validation["worksheet_count"] == 2
    assert validation["parquet_file_count"] == 2
    assert validation["record_count"] == 6
    assert validation["null_count"] == 1


def test_naive_timestamps_without_source_timezone_need_input(tmp_path: Path) -> None:
    source_path = tmp_path / "synthetic.xlsx"
    _write_reference_workbook(source_path)
    service, source = _service_for(source_path, tmp_path / "artifacts")

    started = service.start(ParseRequest(sources=[source]))

    assert started.status == "needs_input"
    status = service.status(started.job_id)
    assert "source_timezone is required" in status.warnings[0]


def test_conflicting_duplicate_point_metadata_need_input(tmp_path: Path) -> None:
    source_path = tmp_path / "conflict.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "负荷A"
    first.append(["time", "DUPLICATE"])
    first.append(["点位描述", "设备A|负荷|kW"])
    first.append([datetime(2026, 1, 1, 8, 0), 1])
    first.append([datetime(2026, 1, 1, 8, 1), 2])
    second = workbook.create_sheet("负荷B")
    second.append(["time", "DUPLICATE"])
    second.append(["点位描述", "设备B|温度|°C"])
    second.append([datetime(2026, 1, 1, 8, 0), 3])
    second.append([datetime(2026, 1, 1, 8, 1), 4])
    workbook.save(source_path)
    service, source = _service_for(source_path, tmp_path / "artifacts")

    started = service.start(
        ParseRequest(
            sources=[source],
            parse_hints={"source_timezone": "Asia/Shanghai"},
        )
    )

    assert started.status == "needs_input"
    assert "conflicting metadata" in service.status(started.job_id).warnings[0]


def test_recoverable_empty_sheet_produces_partial_success(tmp_path: Path) -> None:
    source_path = tmp_path / "partial.xlsx"
    workbook = Workbook()
    good = workbook.active
    good.title = "负荷"
    good.append(["time", "GOOD"])
    good.append(["点位描述", "设备|负荷|kW"])
    good.append([datetime(2026, 1, 1, 8, 0), 1])
    good.append([datetime(2026, 1, 1, 8, 1), 2])
    empty = workbook.create_sheet("空表")
    empty.append(["time", "GHOST"])
    empty.append(["点位描述", "设备|空值|kW"])
    workbook.save(source_path)
    service, source = _service_for(source_path, tmp_path / "artifacts")

    started = service.start(
        ParseRequest(
            sources=[source],
            parse_hints={"source_timezone": "Asia/Shanghai"},
        )
    )

    assert started.status == "succeeded"
    result = service.result(started.job_id)
    assert result.summary.successful_worksheet_count == 1
    assert result.summary.failed_worksheet_count == 1
    bundle_path = (
        tmp_path / "artifacts" / "dataset-bundles" / result.dataset_bundle_id
    )
    points = pq.read_table(bundle_path / "points.parquet")["point_id"].to_pylist()
    assert points == ["GOOD"]
    validation = validate_dataset_bundle(bundle_path)
    assert validation["parquet_file_count"] == 1
    assert validation["failed_worksheet_count"] == 1
